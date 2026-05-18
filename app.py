"""COD 审核系统 — Streamlit Web 应用

支持:
- 上传 Excel/PDF/图片 原始记录文件
- 加载真实样本示例数据
- 执行 31 条 HJ 828-2017 审核规则
- 查看审核报告和审核历史
"""

import streamlit as st
import json
from datetime import date, datetime
from pathlib import Path

from models.record import CODRecord, SampleRow
from models.audit_result import AuditStatus
from engine.auditor import Auditor
from engine.batch_context import build_batch_context
from engine.calculator import compute_fas_concentration, compute_blank_average, recompute_all_cod
from parsers.excel_parser import build_demo_record
from parsers import parse as parse_file

# ============================================================
# 页面配置
# ============================================================
st.set_page_config(
    page_title="COD 原始记录审核系统",
    page_icon="\U0001f9ea",
    layout="wide",
)

# ============================================================
# 状态初始化
# ============================================================
if "record" not in st.session_state:
    st.session_state.record = None
if "report" not in st.session_state:
    st.session_state.report = None
if "parsed_file_name" not in st.session_state:
    st.session_state.parsed_file_name = ""

STATUS_COLORS = {
    "通过": "#27ae60",
    "信息": "#3498db",
    "警告": "#f39c12",
    "不通过": "#e74c3c",
}
STATUS_ICONS = {
    "通过": "✅",
    "信息": "ℹ️",
    "警告": "⚠️",
    "不通过": "❌",
}


# ============================================================
# 侧边栏
# ============================================================
def sidebar():
    st.sidebar.title("\U0001f9ea COD 审核系统")
    st.sidebar.caption("依据 HJ 828-2017 水质 化学需氧量的测定 重铬酸盐法")

    page = st.sidebar.radio(
        "导航",
        ["\U0001f3e0 首页", "\U0001f4e4 上传审核", "\U0001f4ca 审核报告", "\U0001f4c2 审核历史"],
        label_visibility="collapsed",
    )

    st.sidebar.divider()
    st.sidebar.markdown("### 统计概览")
    try:
        from db.repository import get_conn
        conn = get_conn()
        total = conn.execute("SELECT COUNT(*) FROM audit_records").fetchone()[0]
        fail_count = conn.execute(
            "SELECT COUNT(*) FROM audit_records WHERE overall_verdict = '不通过'"
        ).fetchone()[0]
        conn.close()
        st.sidebar.metric("累计审核", f"{total} 份")
        st.sidebar.metric("不通过", f"{fail_count} 份")
    except Exception:
        st.sidebar.caption("数据库未初始化")

    st.sidebar.divider()
    st.sidebar.caption("v1.0 | HJ 828-2017 | 31 条规则")

    return page


# ============================================================
# 首页
# ============================================================
def page_home():
    st.title("\U0001f9ea COD 化学需氧量原始记录审核系统")
    st.caption("依据 HJ 828-2017 水质 化学需氧量的测定 重铬酸盐法 — 31 条审核规则")

    col1, col2 = st.columns(2)

    with col1:
        st.markdown("""
        ### 功能说明

        本系统用于审核 COD 分析原始记录是否符合 HJ 828-2017 标准方法要求。

        **审核规则覆盖 11 个类别：**

        | 类别 | 规则数 | 核心检查项 |
        |------|--------|-----------|
        | 标定审核 | 4 | 日期自洽、平行双样、偏差、浓度验证 |
        | 空白审核 | 3 | 数量、一致性、合理性 |
        | 精密度审核 | 3 | 平行数量、RPD偏差、检出限豁免 |
        | 准确度审核 | 2 | 质控样存在性、保证值范围 |
        | 计算审核 | 2 | COD重算验证、空白均值使用 |
        | 仪器溯源 | 3 | 滴定管、消解装置、天平有效期 |
        | 方法参数 | 5 | 浓度级别、试剂匹配、用量、保存时间 |
        | 氯离子干扰 | 2 | 浓度上限、HgSO₄加入量 |
        | 检出限范围 | 3 | 检出限、测定下限、测定上限 |
        | 稀释审核 | 2 | 稀释范围、取样一致性 |
        | 结果表示 | 2 | 有效数字修约、低于检出限格式 |
        """)

    with col2:
        st.markdown("### 快速开始")

        if st.button("\U0001f680 加载真实样本数据演示", type="primary", use_container_width=True):
            with st.spinner("加载示例数据..."):
                st.session_state.record = build_demo_record()
                st.session_state.report = None
                st.session_state.parsed_file_name = "化学需氧量(容量法)测定原始记录表.xls"
            st.success("示例数据已加载！点击「开始审核」按钮运行审核引擎")
            st.rerun()

        st.divider()

        st.markdown("### 上传原始记录文件")
        uploaded = st.file_uploader(
            "支持 Excel / PDF / DOCX / 图片 格式",
            type=["xls", "xlsx", "pdf", "docx", "png", "jpg", "jpeg"],
            key="home_upload",
        )
        if uploaded is not None:
            _handle_upload(uploaded)

        st.divider()

        st.markdown("### 当前状态")
        if st.session_state.record:
            rec = st.session_state.record
            st.success(f"**已加载**: {st.session_state.parsed_file_name or rec.record_id}")
            st.markdown(f"- 记录编号: `{rec.record_id}`")
            st.markdown(f"- 任务编号: `{rec.task_id}`")
            st.markdown(f"- 样品数: {len(rec.samples)} (含 {len(rec.blanks)} 空白, {len(rec.qc_standards)} 标样)")
            st.markdown(f"- 分析日期: {rec.analysis_date}")

            if st.session_state.report:
                st.info(f"审核结论: **{st.session_state.report.overall_verdict}**")
        else:
            st.info("尚未加载数据，请上传文件或加载示例数据")


# ============================================================
# 上传审核页
# ============================================================
def page_upload():
    st.title("\U0001f4e4 上传 & 审核")

    col1, col2 = st.columns([1, 1])

    with col1:
        st.markdown("### 上传原始记录文件")
        uploaded = st.file_uploader(
            "支持 Excel / PDF / DOCX / 图片 格式",
            type=["xls", "xlsx", "pdf", "docx", "png", "jpg", "jpeg"],
            key="upload_page",
        )
        if uploaded is not None:
            _handle_upload(uploaded)

    with col2:
        st.markdown("### 或使用示例数据")
        if st.button("\U0001f680 加载真实样本数据", type="primary", use_container_width=True):
            st.session_state.record = build_demo_record()
            st.session_state.report = None
            st.session_state.parsed_file_name = "化学需氧量(容量法)测定原始记录表.xls"
            st.rerun()

    st.divider()

    if st.session_state.record:
        st.markdown("### 批次数据概览")
        rec = st.session_state.record
        _render_record_overview(rec)


def _handle_upload(uploaded):
    """处理上传文件 — 多格式支持 + 可选 AI fallback"""
    from parsers.ai_mapper import AIFieldMapper

    with st.spinner(f"正在解析 {uploaded.name}..."):
        result = parse_file(uploaded.getvalue(), uploaded.name)

        if not result.ok:
            st.error(f"文件解析失败: {'; '.join(result.warnings)}")
            st.info("已自动加载示例数据进行演示")
            st.session_state.record = build_demo_record()
            st.session_state.parsed_file_name = f"{uploaded.name} (解析失败，使用示例数据)"
            return

        rec = result.record
        warnings = result.warnings

        # AI fallback for missing fields
        use_ai = st.checkbox(
            "AI 辅助补全缺失字段",
            value=False,
            help="当启发式解析无法提取全部字段时，使用 AI 从原始文本中补全",
            key=f"ai_{uploaded.name}",
        )
        if use_ai and warnings:
            ai_api_key = st.secrets.get("ANTHROPIC_API_KEY", "")
            mapper = AIFieldMapper(api_key=ai_api_key)
            if mapper.available:
                with st.spinner("AI 正在补全缺失字段..."):
                    # Re-extract raw text for AI
                    raw_text = _get_raw_text(uploaded)
                    rec, changed = mapper.fill_gaps(raw_text, rec)
                    if changed:
                        st.success("AI 已补全部分字段")
                    else:
                        st.info("AI 未找到可补全的字段")
            else:
                st.warning("未配置 ANTHROPIC_API_KEY，无法使用 AI 功能")

        st.session_state.record = rec
        st.session_state.report = None
        st.session_state.parsed_file_name = uploaded.name

        msg = f"解析成功！提取到 {len(rec.samples)} 条样品记录（{result.source_format} 格式）"
        if result.used_ai_fallback:
            msg += " | AI 辅助提取"
        if warnings:
            msg += f" | {len(warnings)} 个警告"
        st.success(msg)
        if warnings:
            for w in warnings[:3]:
                st.warning(w)
        st.rerun()


def _get_raw_text(uploaded) -> str:
    """从上传文件获取原始文本 (供 AI mapper 使用)"""
    ext = uploaded.name.rsplit('.', 1)[-1].lower() if '.' in uploaded.name else ''
    file_bytes = uploaded.getvalue()

    if ext == 'pdf':
        import fitz
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        try:
            return '\n'.join(doc[i].get_text("text") for i in range(doc.page_count))
        finally:
            doc.close()
    elif ext in ('xls', 'xlsx'):
        return _read_excel_text(file_bytes, ext)
    elif ext == 'docx':
        from docx import Document
        from io import BytesIO
        doc = Document(BytesIO(file_bytes))
        lines = []
        for table in doc.tables:
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells)
        return '\n'.join(lines)
    elif ext in ('png', 'jpg', 'jpeg'):
        import numpy as np
        from PIL import Image
        from io import BytesIO
        try:
            import easyocr
            reader = easyocr.Reader(['ch_sim', 'en'], gpu=False)
            img = Image.open(BytesIO(file_bytes))
            if img.mode in ('RGBA', 'L'):
                img = img.convert('RGB')
            results = reader.readtext(np.array(img))
            results.sort(key=lambda r: (r[0][0][1], r[0][0][0]))
            return '\n'.join(text for (_bbox, text, _conf) in results if text.strip())
        except ImportError:
            return ""
    return ""


def _read_excel_text(file_bytes: bytes, ext: str) -> str:
    """读取 Excel 文件原始文本"""
    from io import BytesIO
    lines = []
    if ext == 'xls':
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_bytes)
        for sheet in wb.sheets():
            for row in range(sheet.nrows):
                row_text = ' '.join(
                    str(sheet.cell_value(row, col)) for col in range(sheet.ncols)
                )
                lines.append(row_text)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                row_text = ' '.join(str(c) if c is not None else '' for c in row)
                if row_text.strip():
                    lines.append(row_text)
    return '\n'.join(lines)


def _render_record_overview(rec: CODRecord):
    """渲染批次数据概览"""
    col1, col2, col3, col4, col5 = st.columns(5)
    col1.metric("样品总数", len(rec.samples))
    col2.metric("空白数", len(rec.blanks))
    col3.metric("标样数", len(rec.qc_standards))
    col4.metric("平行样对数", len(rec.parallels) // 2)
    col5.metric("方法级别", "低浓度" if rec.overall_method_level == "low" else "高浓度" if rec.overall_method_level == "high" else "混合")

    st.divider()

    # 样品列表
    st.markdown("**样品明细**")
    sample_data = []
    for s in rec.samples:
        sample_data.append({
            "序号": s.seq,
            "样品编号": s.sample_id,
            "类型": "空白" if s.is_blank else "标样" if s.is_qc_standard else "平行" if s.is_parallel else "样品",
            "取样体积 (ml)": s.volume,
            "稀释倍数": f"{s.dilution_factor:.0f}" if s.dilution_factor > 1 else "-",
            "净用量 (ml)": f"{s.net_volume:.2f}" if s.net_volume > 0 else "-",
            "填报值": s.cod_display,
            "Cl⁻ (mg/L)": f"{s.cl_estimate:.0f}" if s.cl_estimate else "-",
            "HgSO₄ (ml)": f"{s.hgso4_added:.2f}" if s.hgso4_added else "-",
        })
    st.dataframe(sample_data, use_container_width=True, hide_index=True)

    # 标定数据
    if rec.fas_std.volumes:
        st.markdown("**FAS 标定数据**")
        c1, c2, c3 = st.columns(3)
        c1.metric("重铬酸钾浓度", f"{rec.fas_std.k2cr2o7_conc:.4f} mol/L")
        c2.metric("标定体积", f"{rec.fas_std.average_volume:.2f} ml ({rec.fas_std.volumes[0]:.2f}, {rec.fas_std.volumes[1]:.2f})")
        c3.metric("标定浓度", f"{rec.fas_std.calculated_conc:.6f} mol/L")

    # 仪器
    if rec.instruments:
        st.markdown("**仪器溯源**")
        inst_data = [{
            "仪器名称": i.name,
            "编号": i.serial_no,
            "溯源有效期": str(i.calibration_expiry) if i.calibration_expiry else "-",
            "溯源方式": i.calibration_method or "-",
        } for i in rec.instruments]
        st.dataframe(inst_data, use_container_width=True, hide_index=True)

    # 审核按钮
    st.divider()
    c1, c2, c3 = st.columns([1, 2, 1])
    with c2:
        if st.button("\U0001f50d 开始审核", type="primary", use_container_width=True):
            with st.spinner("正在执行 31 条审核规则..."):
                auditor = Auditor()
                st.session_state.report = auditor.audit(st.session_state.record)
            if st.session_state.report.overall_pass:
                st.balloons()
            st.rerun()


# ============================================================
# 审核报告页
# ============================================================
def page_report():
    st.title("\U0001f4ca 审核报告")

    if not st.session_state.record or not st.session_state.report:
        st.info("ℹ️ 请先上传原始记录文件或加载示例数据，然后执行审核")
        if st.button("← 前往上传页面"):
            st.switch_page("app.py")  # fallback nav
        return

    rec = st.session_state.record
    report = st.session_state.report

    _render_report_header(rec, report)
    st.divider()
    _render_report_items(report)
    st.divider()
    _render_report_details(rec, report)


def _render_report_header(rec: CODRecord, report):
    """渲染报告头部"""
    # 结论横幅
    verdict = report.overall_verdict
    color_map = {"通过": "#27ae60", "有条件通过": "#f39c12", "不通过": "#e74c3c"}
    bg = color_map.get(verdict, "#95a5a6")

    st.markdown(f"""
    <div style="background:{bg}; padding:20px; border-radius:10px; color:white; text-align:center; margin-bottom:20px">
        <h1 style="margin:0; font-size:2.5em">{verdict}</h1>
        <p style="margin:5px 0 0 0; opacity:0.9">审核结论</p>
    </div>
    """, unsafe_allow_html=True)

    # 统计卡片
    col1, col2, col3, col4, col5 = st.columns(5)
    col1.metric("✅ 通过", report.pass_count)
    col2.metric("ℹ️ 信息", report.info_count)
    col3.metric("⚠️ 警告", report.warning_count)
    col4.metric("❌ 不通过", report.fail_count)
    col5.metric("\U0001f4cb 总规则", len(report.items))

    # 批次信息
    st.markdown(f"""
    | 项目 | 内容 |
    |------|------|
    | 记录编号 | `{rec.record_id}` |
    | 任务编号 | `{rec.task_id}` |
    | 分析日期 | {rec.analysis_date} |
    | 分析方法 | {rec.method_ref[:50]}... |
    | 审核时间 | {report.audit_time[:19]} |
    """)


def _render_report_items(report):
    """渲染每条审核结果"""
    st.markdown("### 审核结果明细")

    # 分类折叠展示
    categories = {}
    for item in report.items:
        cat = item.category or "其他"
        if cat not in categories:
            categories[cat] = []
        categories[cat].append(item)

    for cat, items in categories.items():
        pass_count = sum(1 for i in items if i.status == AuditStatus.PASS)
        fail_count = sum(1 for i in items if i.status == AuditStatus.FAIL)
        warn_count = sum(1 for i in items if i.status == AuditStatus.WARNING)
        info_count = sum(1 for i in items if i.status == AuditStatus.INFO)

        status_badges = []
        if pass_count: status_badges.append(f"✅ {pass_count}")
        if info_count: status_badges.append(f"ℹ️ {info_count}")
        if warn_count: status_badges.append(f"⚠️ {warn_count}")
        if fail_count: status_badges.append(f"❌ {fail_count}")

        with st.expander(f"{cat}  {' '.join(status_badges)}", expanded=(fail_count > 0)):
            for item in items:
                icon = STATUS_ICONS.get(item.status.value, "❓")
                color = STATUS_COLORS.get(item.status.value, "#95a5a6")

                st.markdown(f"""
                <div style="border-left:4px solid {color}; padding:8px 12px; margin:4px 0; background:#f8f9fa; border-radius:4px">
                    <strong>{icon} [{item.code}] {item.name}</strong>
                    <span style="color:{color}; float:right; font-weight:bold">{item.status.value}</span>
                    <br><small>{item.detail}</small>
                </div>
                """, unsafe_allow_html=True)

                if item.suggestion:
                    st.caption(f"  \U0001f4a1 {item.suggestion}")


def _render_report_details(rec: CODRecord, report):
    """渲染计算详情"""
    st.markdown("### 计算详情")

    tab1, tab2 = st.tabs(["COD 重算结果", "原始数据"])

    with tab1:
        recompute_all_cod(rec)
        calc_data = []
        for s in rec.samples:
            calc_data.append({
                "样品编号": s.sample_id,
                "类型": "空白" if s.is_blank else "标样" if s.is_qc_standard else "样品",
                "净用量 (ml)": f"{s.net_volume:.2f}",
                "系统重算 (mg/L)": f"{s.cod_calculated:.2f}" if s.cod_calculated else "-",
                "填报值": s.cod_display,
            })
        st.dataframe(calc_data, use_container_width=True, hide_index=True)

        blank_avg = compute_blank_average(rec)
        fas_conc = compute_fas_concentration(rec)
        c1, c2 = st.columns(2)
        c1.metric("空白均值", f"{blank_avg:.3f} ml")
        c2.metric("FAS 浓度", f"{fas_conc:.6f} mol/L")

    with tab2:
        st.json(json.dumps(rec.to_dict(), ensure_ascii=False, default=str))


# ============================================================
# 审核历史页
# ============================================================
def page_history():
    st.title("\U0001f4c2 审核历史")

    try:
        from db.repository import list_records, load_audit
        records = list_records(limit=100)
    except Exception as e:
        st.warning(f"数据库连接失败: {e}")
        records = []

    if not records:
        st.info("暂无审核记录")
        return

    st.markdown(f"共 {len(records)} 条审核记录")

    for r in records:
        verdict = r.get("overall_verdict", "通过")
        color = {"通过": "#27ae60", "有条件通过": "#f39c12", "不通过": "#e74c3c"}.get(verdict, "#95a5a6")

        with st.expander(
            f"[{r.get('analysis_date', '-')}] {r.get('id', '-')} — {verdict} "
            f"({r.get('pass_count', 0)}✅ {r.get('fail_count', 0)}❌)"
        ):
            col1, col2, col3 = st.columns(3)
            col1.markdown(f"**记录编号**: `{r.get('id')}`")
            col1.markdown(f"**任务编号**: `{r.get('task_id', '-')}`")
            col2.markdown(f"**分析日期**: {r.get('analysis_date', '-')}")
            col2.markdown(f"**方法**: {r.get('method_ref', '-')[:60]}")
            col3.metric("通过", r.get('pass_count', 0))
            col3.metric("不通过", r.get('fail_count', 0))

            if st.button(f"查看详情", key=f"detail_{r.get('id')}"):
                record, report = load_audit(r.get('id'))
                if report:
                    st.session_state.record = record
                    st.session_state.report = report
                    st.rerun()


# ============================================================
# 主入口
# ============================================================
def main():
    page = sidebar()

    if page == "\U0001f3e0 首页":
        page_home()
    elif page == "\U0001f4e4 上传审核":
        page_upload()
    elif page == "\U0001f4ca 审核报告":
        page_report()
    elif page == "\U0001f4c2 审核历史":
        page_history()


if __name__ == "__main__":
    main()
